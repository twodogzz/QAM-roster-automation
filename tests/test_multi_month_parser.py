from __future__ import annotations

from pathlib import Path
import tempfile
import unittest

from docx import Document

from modules.roster_parser import parse_roster_docx
from modules.sync_service import MODE_DELETE_CURRENT, MODE_FULL, SyncOptions, prepare_sync_plan


class MultiMonthRosterTests(unittest.TestCase):
    def test_parse_roster_docx_supports_next_month_cells(self) -> None:
        docx_path = _build_roster_docx(
            "Queensland Air Museum - Front Counter Roster March 2026 v9",
            [
                ["Sunday", "Monday", "Tuesday"],
                ["30. Someone Else", "31. Wayne Freestun", "1. April Wayne Freestun"],
            ],
        )

        parsed = parse_roster_docx(docx_path)

        self.assertEqual((parsed.year, parsed.month), (2026, 3))
        self.assertEqual(parsed.covered_months, [(2026, 3), (2026, 4)])
        self.assertEqual(
            [event.event_date.isoformat() for event in parsed.roster_events],
            ["2026-03-31", "2026-04-01"],
        )
        self.assertTrue(all(event.summary.endswith("Mar v9") for event in parsed.roster_events))

    def test_parse_roster_docx_supports_previous_month_cells(self) -> None:
        docx_path = _build_roster_docx(
            "Queensland Air Museum - Front Counter Roster January 2026 v4",
            [
                ["Sunday", "Monday"],
                ["31. December Wayne Freestun", "1. Wayne Freestun"],
            ],
        )

        parsed = parse_roster_docx(docx_path)

        self.assertEqual(parsed.covered_months, [(2025, 12), (2026, 1)])
        self.assertEqual(
            [event.event_date.isoformat() for event in parsed.roster_events],
            ["2025-12-31", "2026-01-01"],
        )
        self.assertTrue(all(event.summary.endswith("Jan v4") for event in parsed.roster_events))

    def test_prepare_sync_plan_scopes_deletes_across_all_parsed_months(self) -> None:
        docx_path = _build_roster_docx(
            "Queensland Air Museum - Front Counter Roster March 2026 v9",
            [
                ["Sunday", "Monday", "Tuesday"],
                ["30. Wayne Freestun", "31. Wayne Freestun", "1. April Wayne Freestun"],
            ],
        )

        listed_events = [
            _calendar_event("keep-old-feb", "2026-02-28T09:00:00", "Wayne volunteer Queensland Air Museum - Mar v8"),
            _calendar_event("delete-old-mar", "2026-03-31T09:00:00", "Wayne volunteer Queensland Air Museum - Mar v8"),
            _calendar_event("delete-old-apr", "2026-04-01T09:00:00", "Wayne volunteer Queensland Air Museum - Mar v8"),
            _calendar_event("delete-current-mar", "2026-03-30T09:00:00", "Wayne volunteer Queensland Air Museum - Mar v9"),
        ]

        from modules import sync_service

        original_list_events = sync_service.calendar_api.list_events
        original_set_calendar_id = sync_service.calendar_api.set_calendar_id
        try:
            sync_service.calendar_api.list_events = lambda query_text: listed_events
            sync_service.calendar_api.set_calendar_id = lambda calendar_id: None

            plan = prepare_sync_plan(
                SyncOptions(
                    docx_file=docx_path,
                    mode=MODE_FULL,
                    calendar_id="test-calendar",
                    replace_current=True,
                )
            )
            delete_current_plan = prepare_sync_plan(
                SyncOptions(
                    docx_file=docx_path,
                    mode=MODE_DELETE_CURRENT,
                    calendar_id="test-calendar",
                )
            )
        finally:
            sync_service.calendar_api.list_events = original_list_events
            sync_service.calendar_api.set_calendar_id = original_set_calendar_id

        self.assertEqual(plan.covered_months, [(2026, 3), (2026, 4)])
        self.assertEqual([row["id"] for row in plan.delete_events], ["delete-old-mar", "delete-old-apr", "delete-current-mar"])
        self.assertEqual(len(plan.create_payloads), 3)
        self.assertEqual([row["id"] for row in delete_current_plan.delete_events], ["delete-current-mar"])


def _build_roster_docx(title: str, rows: list[list[str]]) -> str:
    document = Document()
    table = document.add_table(rows=len(rows) + 1, cols=len(rows[0]))
    for col in range(len(rows[0])):
        table.cell(0, col).text = title
    for row_index, row in enumerate(rows, start=1):
        for col_index, value in enumerate(row):
            table.cell(row_index, col_index).text = value

    temp_dir = tempfile.mkdtemp(prefix="qam-roster-test-")
    path = Path(temp_dir) / f"{title.replace(' ', '_')}.docx"
    document.save(path)
    return str(path)


def _calendar_event(event_id: str, start: str, summary: str) -> dict[str, object]:
    return {
        "id": event_id,
        "summary": summary,
        "start": {"dateTime": start},
    }


if __name__ == "__main__":
    unittest.main()
